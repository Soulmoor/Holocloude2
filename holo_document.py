#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
HoloDocument v1.0 - Dokument-Analyse-System

FEATURES:
- PDF strukturiert lesen
- Word-Dokumente (DOCX) parsen
- Markdown analysieren
- Text-Extraktion mit Struktur
- Metadaten auslesen
- Tabellen erkennen

Author: Kira & Claude
Version: 1.0
"""

import os
import re
import hashlib
import logging
from typing import Dict, List, Optional, Tuple, Any
from dataclasses import dataclass, field
from datetime import datetime
from enum import Enum, auto
from pathlib import Path

logger = logging.getLogger("HoloDocument")


# =============================================================================
# OPTIONAL IMPORTS
# =============================================================================

_HAS_PYPDF = False
_HAS_DOCX = False
_HAS_MARKDOWN = False
_HAS_PDFPLUMBER = False

try:
    import PyPDF2
    _HAS_PYPDF = True
    logger.info("PyPDF2 verfuegbar")
except ImportError:
    logger.warning("PyPDF2 nicht verfuegbar")

try:
    import pdfplumber
    _HAS_PDFPLUMBER = True
    logger.info("pdfplumber verfuegbar")
except ImportError:
    logger.warning("pdfplumber nicht verfuegbar")

try:
    from docx import Document as DocxDocument
    from docx.table import Table
    _HAS_DOCX = True
    logger.info("python-docx verfuegbar")
except ImportError:
    logger.warning("python-docx nicht verfuegbar")

try:
    import markdown
    _HAS_MARKDOWN = True
    logger.info("markdown verfuegbar")
except ImportError:
    logger.warning("markdown nicht verfuegbar")


# =============================================================================
# ENUMS UND DATACLASSES
# =============================================================================

class DocumentType(Enum):
    """Dokument-Typen"""
    PDF = "pdf"
    DOCX = "docx"
    DOC = "doc"
    MARKDOWN = "markdown"
    TEXT = "text"
    HTML = "html"
    UNKNOWN = "unbekannt"


class ElementType(Enum):
    """Struktur-Element-Typen"""
    TITLE = "titel"
    HEADING_1 = "ueberschrift_1"
    HEADING_2 = "ueberschrift_2"
    HEADING_3 = "ueberschrift_3"
    PARAGRAPH = "absatz"
    LIST_ITEM = "listenelement"
    TABLE = "tabelle"
    IMAGE = "bild"
    CODE_BLOCK = "codeblock"
    QUOTE = "zitat"
    FOOTNOTE = "fussnote"
    LINK = "link"


@dataclass
class DocumentMetadata:
    """Metadaten eines Dokuments"""
    doc_id: str
    file_path: str
    doc_type: DocumentType
    title: Optional[str]
    author: Optional[str]
    created_date: Optional[str]
    modified_date: Optional[str]
    page_count: int
    word_count: int
    character_count: int
    file_size_bytes: int
    language: Optional[str]


@dataclass
class StructureElement:
    """Ein Struktur-Element im Dokument"""
    element_id: str
    element_type: ElementType
    content: str
    page: int  # Seitennummer (1-basiert)
    position: int  # Position im Dokument
    level: int  # Hierarchie-Level (fuer Ueberschriften)
    style: Optional[str]  # Original-Style (z.B. "Heading 1")


@dataclass
class TableData:
    """Eine extrahierte Tabelle"""
    table_id: str
    page: int
    rows: int
    columns: int
    headers: List[str]
    data: List[List[str]]  # Zeilen x Spalten
    has_header: bool


@dataclass
class DocumentStructure:
    """Strukturierte Darstellung des Dokuments"""
    doc_id: str
    elements: List[StructureElement]
    headings: List[StructureElement]
    tables: List[TableData]
    images: List[Dict[str, Any]]
    links: List[Dict[str, str]]
    toc: List[Dict[str, Any]]  # Table of Contents


@dataclass
class DocumentAnalysisResult:
    """Vollstaendiges Ergebnis einer Dokument-Analyse"""
    metadata: DocumentMetadata
    structure: DocumentStructure
    full_text: str
    markdown_content: str  # Konvertiert zu Markdown
    summary_sentences: List[str]  # Erste/wichtige Saetze
    timestamp: str = field(default_factory=lambda: datetime.now().isoformat())


# =============================================================================
# HOLODOCUMENT
# =============================================================================

class HoloDocument:
    """
    Dokument-Analyse-System.

    Features:
    - PDF/DOCX/Markdown parsen
    - Struktur extrahieren
    - Tabellen erkennen
    - Metadaten auslesen
    - Konvertierung zu Markdown
    """

    def __init__(self, output_dir: str = "data/documents"):
        self.output_dir = Path(output_dir)
        self.output_dir.mkdir(parents=True, exist_ok=True)

        self.has_pypdf = _HAS_PYPDF
        self.has_docx = _HAS_DOCX
        self.has_markdown = _HAS_MARKDOWN
        self.has_pdfplumber = _HAS_PDFPLUMBER

        logger.info("HoloDocument initialisiert")

    # =========================================================================
    # HILFSFUNKTIONEN
    # =========================================================================

    def _generate_id(self, content: str) -> str:
        """Generiert eindeutige ID"""
        return hashlib.md5(content.encode()).hexdigest()[:12]

    def _detect_document_type(self, file_path: str) -> DocumentType:
        """Erkennt den Dokument-Typ"""
        ext = Path(file_path).suffix.lower()

        type_map = {
            '.pdf': DocumentType.PDF,
            '.docx': DocumentType.DOCX,
            '.doc': DocumentType.DOC,
            '.md': DocumentType.MARKDOWN,
            '.markdown': DocumentType.MARKDOWN,
            '.txt': DocumentType.TEXT,
            '.html': DocumentType.HTML,
            '.htm': DocumentType.HTML,
        }

        return type_map.get(ext, DocumentType.UNKNOWN)

    def _count_words(self, text: str) -> int:
        """Zaehlt Woerter im Text"""
        words = re.findall(r'\b\w+\b', text)
        return len(words)

    def _extract_sentences(self, text: str, max_sentences: int = 5) -> List[str]:
        """Extrahiert die ersten/wichtigsten Saetze"""
        sentences = re.split(r'[.!?]+', text)
        sentences = [s.strip() for s in sentences if len(s.strip()) > 20]
        return sentences[:max_sentences]

    def _clean_text(self, text: str) -> str:
        """Bereinigt extrahierten Text"""
        # Mehrfache Leerzeichen/Zeilenumbrueche
        text = re.sub(r'\n{3,}', '\n\n', text)
        text = re.sub(r' {2,}', ' ', text)
        text = re.sub(r'\t+', ' ', text)
        return text.strip()

    # =========================================================================
    # PDF-PARSING
    # =========================================================================

    def _parse_pdf(self, file_path: str) -> Tuple[DocumentMetadata, DocumentStructure, str]:
        """Parst ein PDF-Dokument"""
        doc_id = self._generate_id(file_path)
        file_size = os.path.getsize(file_path)

        elements = []
        tables = []
        full_text = ""
        page_count = 0
        title = None
        author = None
        created = None
        modified = None

        # Mit PyPDF2 Text und Metadaten extrahieren
        if self.has_pypdf:
            try:
                with open(file_path, 'rb') as f:
                    reader = PyPDF2.PdfReader(f)
                    page_count = len(reader.pages)

                    # Metadaten
                    if reader.metadata:
                        title = reader.metadata.get('/Title')
                        author = reader.metadata.get('/Author')
                        created = reader.metadata.get('/CreationDate')
                        modified = reader.metadata.get('/ModDate')

                    # Text pro Seite
                    for page_num, page in enumerate(reader.pages, 1):
                        text = page.extract_text() or ""
                        full_text += text + "\n"

                        # Einfache Absatz-Erkennung
                        paragraphs = text.split('\n\n')
                        for i, para in enumerate(paragraphs):
                            para = para.strip()
                            if not para:
                                continue

                            # Ueberschrift erkennen (kurz, evtl. GROSSBUCHSTABEN)
                            if len(para) < 100 and para.isupper():
                                elem_type = ElementType.HEADING_1
                                level = 1
                            elif len(para) < 80 and para.istitle():
                                elem_type = ElementType.HEADING_2
                                level = 2
                            else:
                                elem_type = ElementType.PARAGRAPH
                                level = 0

                            element = StructureElement(
                                element_id=f"{doc_id}_p{page_num}_e{i}",
                                element_type=elem_type,
                                content=para,
                                page=page_num,
                                position=len(elements),
                                level=level,
                                style=None
                            )
                            elements.append(element)

            except Exception as e:
                logger.error(f"PyPDF2 Fehler: {e}")

        # Mit pdfplumber Tabellen extrahieren
        if self.has_pdfplumber:
            try:
                with pdfplumber.open(file_path) as pdf:
                    for page_num, page in enumerate(pdf.pages, 1):
                        page_tables = page.extract_tables()

                        for t_idx, table_data in enumerate(page_tables):
                            if not table_data:
                                continue

                            # Headers (erste Zeile)
                            headers = [str(c) if c else "" for c in table_data[0]]
                            data = [[str(c) if c else "" for c in row] for row in table_data[1:]]

                            table = TableData(
                                table_id=f"{doc_id}_t{page_num}_{t_idx}",
                                page=page_num,
                                rows=len(table_data),
                                columns=len(headers),
                                headers=headers,
                                data=data,
                                has_header=True
                            )
                            tables.append(table)

            except Exception as e:
                logger.error(f"pdfplumber Fehler: {e}")

        # Bereinigen
        full_text = self._clean_text(full_text)

        # Metadaten
        metadata = DocumentMetadata(
            doc_id=doc_id,
            file_path=file_path,
            doc_type=DocumentType.PDF,
            title=title,
            author=author,
            created_date=created,
            modified_date=modified,
            page_count=page_count,
            word_count=self._count_words(full_text),
            character_count=len(full_text),
            file_size_bytes=file_size,
            language=None
        )

        # Struktur
        headings = [e for e in elements if e.element_type in
                    [ElementType.HEADING_1, ElementType.HEADING_2, ElementType.HEADING_3]]

        # TOC aus Ueberschriften generieren
        toc = [{"level": h.level, "title": h.content, "page": h.page} for h in headings]

        structure = DocumentStructure(
            doc_id=doc_id,
            elements=elements,
            headings=headings,
            tables=tables,
            images=[],  # PDF-Bilder sind komplex zu extrahieren
            links=[],
            toc=toc
        )

        return metadata, structure, full_text

    # =========================================================================
    # DOCX-PARSING
    # =========================================================================

    def _parse_docx(self, file_path: str) -> Tuple[DocumentMetadata, DocumentStructure, str]:
        """Parst ein DOCX-Dokument"""
        doc_id = self._generate_id(file_path)
        file_size = os.path.getsize(file_path)

        elements = []
        tables = []
        links = []
        full_text = ""

        if not self.has_docx:
            # Fallback: Kein python-docx
            metadata = DocumentMetadata(
                doc_id=doc_id,
                file_path=file_path,
                doc_type=DocumentType.DOCX,
                title=None, author=None, created_date=None, modified_date=None,
                page_count=0, word_count=0, character_count=0,
                file_size_bytes=file_size, language=None
            )
            structure = DocumentStructure(
                doc_id=doc_id, elements=[], headings=[], tables=[],
                images=[], links=[], toc=[]
            )
            return metadata, structure, ""

        try:
            doc = DocxDocument(file_path)

            # Core Properties
            title = doc.core_properties.title
            author = doc.core_properties.author
            created = str(doc.core_properties.created) if doc.core_properties.created else None
            modified = str(doc.core_properties.modified) if doc.core_properties.modified else None

            # Paragraphen durchgehen
            for i, para in enumerate(doc.paragraphs):
                text = para.text.strip()
                if not text:
                    continue

                full_text += text + "\n"

                # Style analysieren
                style_name = para.style.name if para.style else ""

                if "Heading 1" in style_name or "Ueberschrift 1" in style_name:
                    elem_type = ElementType.HEADING_1
                    level = 1
                elif "Heading 2" in style_name or "Ueberschrift 2" in style_name:
                    elem_type = ElementType.HEADING_2
                    level = 2
                elif "Heading 3" in style_name or "Ueberschrift 3" in style_name:
                    elem_type = ElementType.HEADING_3
                    level = 3
                elif "Title" in style_name or "Titel" in style_name:
                    elem_type = ElementType.TITLE
                    level = 0
                elif "List" in style_name:
                    elem_type = ElementType.LIST_ITEM
                    level = 0
                elif "Quote" in style_name or "Zitat" in style_name:
                    elem_type = ElementType.QUOTE
                    level = 0
                else:
                    elem_type = ElementType.PARAGRAPH
                    level = 0

                element = StructureElement(
                    element_id=f"{doc_id}_e{i}",
                    element_type=elem_type,
                    content=text,
                    page=1,  # DOCX hat keine exakten Seitenzahlen
                    position=i,
                    level=level,
                    style=style_name
                )
                elements.append(element)

                # Links extrahieren
                for run in para.runs:
                    if run.font.underline and 'http' in text.lower():
                        # Vereinfachte Link-Erkennung
                        url_match = re.search(r'https?://\S+', text)
                        if url_match:
                            links.append({"text": text, "url": url_match.group()})

            # Tabellen
            for t_idx, table in enumerate(doc.tables):
                rows = []
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells]
                    rows.append(cells)

                if rows:
                    headers = rows[0]
                    data = rows[1:] if len(rows) > 1 else []

                    table_data = TableData(
                        table_id=f"{doc_id}_t{t_idx}",
                        page=1,
                        rows=len(rows),
                        columns=len(headers),
                        headers=headers,
                        data=data,
                        has_header=True
                    )
                    tables.append(table_data)

        except Exception as e:
            logger.error(f"DOCX Fehler: {e}")
            title = author = created = modified = None

        full_text = self._clean_text(full_text)

        # Metadaten
        metadata = DocumentMetadata(
            doc_id=doc_id,
            file_path=file_path,
            doc_type=DocumentType.DOCX,
            title=title,
            author=author,
            created_date=created,
            modified_date=modified,
            page_count=max(1, len(full_text) // 3000),  # Geschaetzt
            word_count=self._count_words(full_text),
            character_count=len(full_text),
            file_size_bytes=file_size,
            language=None
        )

        # Struktur
        headings = [e for e in elements if e.element_type in
                    [ElementType.TITLE, ElementType.HEADING_1,
                     ElementType.HEADING_2, ElementType.HEADING_3]]

        toc = [{"level": h.level, "title": h.content, "page": 1} for h in headings]

        structure = DocumentStructure(
            doc_id=doc_id,
            elements=elements,
            headings=headings,
            tables=tables,
            images=[],
            links=links,
            toc=toc
        )

        return metadata, structure, full_text

    # =========================================================================
    # MARKDOWN-PARSING
    # =========================================================================

    def _parse_markdown(self, file_path: str) -> Tuple[DocumentMetadata, DocumentStructure, str]:
        """Parst ein Markdown-Dokument"""
        doc_id = self._generate_id(file_path)
        file_size = os.path.getsize(file_path)

        with open(file_path, 'r', encoding='utf-8') as f:
            content = f.read()

        elements = []
        tables = []
        links = []
        images = []

        lines = content.split('\n')
        full_text = ""
        current_para = []

        for i, line in enumerate(lines):
            # Ueberschriften
            if line.startswith('# '):
                if current_para:
                    full_text += ' '.join(current_para) + '\n'
                    current_para = []

                element = StructureElement(
                    element_id=f"{doc_id}_e{len(elements)}",
                    element_type=ElementType.HEADING_1,
                    content=line[2:].strip(),
                    page=1,
                    position=len(elements),
                    level=1,
                    style="h1"
                )
                elements.append(element)
                full_text += line[2:].strip() + '\n'

            elif line.startswith('## '):
                if current_para:
                    full_text += ' '.join(current_para) + '\n'
                    current_para = []

                element = StructureElement(
                    element_id=f"{doc_id}_e{len(elements)}",
                    element_type=ElementType.HEADING_2,
                    content=line[3:].strip(),
                    page=1,
                    position=len(elements),
                    level=2,
                    style="h2"
                )
                elements.append(element)
                full_text += line[3:].strip() + '\n'

            elif line.startswith('### '):
                if current_para:
                    full_text += ' '.join(current_para) + '\n'
                    current_para = []

                element = StructureElement(
                    element_id=f"{doc_id}_e{len(elements)}",
                    element_type=ElementType.HEADING_3,
                    content=line[4:].strip(),
                    page=1,
                    position=len(elements),
                    level=3,
                    style="h3"
                )
                elements.append(element)
                full_text += line[4:].strip() + '\n'

            # Listen
            elif re.match(r'^[\*\-\+]\s', line):
                element = StructureElement(
                    element_id=f"{doc_id}_e{len(elements)}",
                    element_type=ElementType.LIST_ITEM,
                    content=line[2:].strip(),
                    page=1,
                    position=len(elements),
                    level=0,
                    style="list"
                )
                elements.append(element)
                full_text += line[2:].strip() + '\n'

            # Code-Bloecke
            elif line.startswith('```'):
                element = StructureElement(
                    element_id=f"{doc_id}_e{len(elements)}",
                    element_type=ElementType.CODE_BLOCK,
                    content=line,
                    page=1,
                    position=len(elements),
                    level=0,
                    style="code"
                )
                elements.append(element)

            # Zitate
            elif line.startswith('> '):
                element = StructureElement(
                    element_id=f"{doc_id}_e{len(elements)}",
                    element_type=ElementType.QUOTE,
                    content=line[2:].strip(),
                    page=1,
                    position=len(elements),
                    level=0,
                    style="quote"
                )
                elements.append(element)
                full_text += line[2:].strip() + '\n'

            # Links extrahieren
            elif '[' in line and '](' in line:
                link_matches = re.findall(r'\[([^\]]+)\]\(([^)]+)\)', line)
                for text, url in link_matches:
                    if url.startswith('http'):
                        links.append({"text": text, "url": url})

                # Auch als Absatz
                current_para.append(line)

            # Bilder
            elif line.startswith('!['):
                img_match = re.match(r'!\[([^\]]*)\]\(([^)]+)\)', line)
                if img_match:
                    images.append({
                        "alt": img_match.group(1),
                        "src": img_match.group(2)
                    })

            # Leerzeile = Absatz-Ende
            elif not line.strip():
                if current_para:
                    para_text = ' '.join(current_para)
                    element = StructureElement(
                        element_id=f"{doc_id}_e{len(elements)}",
                        element_type=ElementType.PARAGRAPH,
                        content=para_text,
                        page=1,
                        position=len(elements),
                        level=0,
                        style="p"
                    )
                    elements.append(element)
                    full_text += para_text + '\n'
                    current_para = []

            # Normaler Text
            else:
                current_para.append(line)

        # Letzter Absatz
        if current_para:
            para_text = ' '.join(current_para)
            element = StructureElement(
                element_id=f"{doc_id}_e{len(elements)}",
                element_type=ElementType.PARAGRAPH,
                content=para_text,
                page=1,
                position=len(elements),
                level=0,
                style="p"
            )
            elements.append(element)
            full_text += para_text + '\n'

        full_text = self._clean_text(full_text)

        # Titel aus erster H1
        title = None
        for e in elements:
            if e.element_type == ElementType.HEADING_1:
                title = e.content
                break

        # Metadaten
        metadata = DocumentMetadata(
            doc_id=doc_id,
            file_path=file_path,
            doc_type=DocumentType.MARKDOWN,
            title=title,
            author=None,
            created_date=None,
            modified_date=None,
            page_count=1,
            word_count=self._count_words(full_text),
            character_count=len(full_text),
            file_size_bytes=file_size,
            language=None
        )

        # Struktur
        headings = [e for e in elements if e.element_type in
                    [ElementType.HEADING_1, ElementType.HEADING_2, ElementType.HEADING_3]]

        toc = [{"level": h.level, "title": h.content, "page": 1} for h in headings]

        structure = DocumentStructure(
            doc_id=doc_id,
            elements=elements,
            headings=headings,
            tables=tables,
            images=images,
            links=links,
            toc=toc
        )

        return metadata, structure, full_text

    # =========================================================================
    # TEXT-PARSING (Plaintext)
    # =========================================================================

    def _parse_text(self, file_path: str) -> Tuple[DocumentMetadata, DocumentStructure, str]:
        """Parst eine Text-Datei"""
        doc_id = self._generate_id(file_path)
        file_size = os.path.getsize(file_path)

        with open(file_path, 'r', encoding='utf-8', errors='ignore') as f:
            content = f.read()

        full_text = self._clean_text(content)

        # Einfache Absatz-Segmentierung
        paragraphs = full_text.split('\n\n')
        elements = []

        for i, para in enumerate(paragraphs):
            para = para.strip()
            if not para:
                continue

            element = StructureElement(
                element_id=f"{doc_id}_e{i}",
                element_type=ElementType.PARAGRAPH,
                content=para,
                page=1,
                position=i,
                level=0,
                style=None
            )
            elements.append(element)

        metadata = DocumentMetadata(
            doc_id=doc_id,
            file_path=file_path,
            doc_type=DocumentType.TEXT,
            title=Path(file_path).stem,
            author=None,
            created_date=None,
            modified_date=None,
            page_count=1,
            word_count=self._count_words(full_text),
            character_count=len(full_text),
            file_size_bytes=file_size,
            language=None
        )

        structure = DocumentStructure(
            doc_id=doc_id,
            elements=elements,
            headings=[],
            tables=[],
            images=[],
            links=[],
            toc=[]
        )

        return metadata, structure, full_text

    # =========================================================================
    # HAUPTFUNKTIONEN
    # =========================================================================

    def analyze_document(self, file_path: str) -> Optional[DocumentAnalysisResult]:
        """
        Analysiert ein Dokument vollstaendig.

        Unterstuetzt: PDF, DOCX, Markdown, Text
        """
        if not os.path.exists(file_path):
            logger.error(f"Datei nicht gefunden: {file_path}")
            return None

        doc_type = self._detect_document_type(file_path)

        # Parsen
        if doc_type == DocumentType.PDF:
            metadata, structure, full_text = self._parse_pdf(file_path)
        elif doc_type == DocumentType.DOCX:
            metadata, structure, full_text = self._parse_docx(file_path)
        elif doc_type == DocumentType.MARKDOWN:
            metadata, structure, full_text = self._parse_markdown(file_path)
        elif doc_type == DocumentType.TEXT:
            metadata, structure, full_text = self._parse_text(file_path)
        else:
            logger.warning(f"Nicht unterstuetzter Dokumenttyp: {doc_type}")
            return None

        # Zu Markdown konvertieren
        markdown_content = self._convert_to_markdown(structure)

        # Zusammenfassung
        summary_sentences = self._extract_sentences(full_text, 5)

        return DocumentAnalysisResult(
            metadata=metadata,
            structure=structure,
            full_text=full_text,
            markdown_content=markdown_content,
            summary_sentences=summary_sentences
        )

    def _convert_to_markdown(self, structure: DocumentStructure) -> str:
        """Konvertiert die Dokumentstruktur zu Markdown"""
        lines = []

        for elem in structure.elements:
            if elem.element_type == ElementType.TITLE:
                lines.append(f"# {elem.content}\n")
            elif elem.element_type == ElementType.HEADING_1:
                lines.append(f"# {elem.content}\n")
            elif elem.element_type == ElementType.HEADING_2:
                lines.append(f"## {elem.content}\n")
            elif elem.element_type == ElementType.HEADING_3:
                lines.append(f"### {elem.content}\n")
            elif elem.element_type == ElementType.LIST_ITEM:
                lines.append(f"- {elem.content}")
            elif elem.element_type == ElementType.QUOTE:
                lines.append(f"> {elem.content}\n")
            elif elem.element_type == ElementType.CODE_BLOCK:
                lines.append(f"```\n{elem.content}\n```\n")
            elif elem.element_type == ElementType.PARAGRAPH:
                lines.append(f"{elem.content}\n")

        # Tabellen
        for table in structure.tables:
            lines.append("")
            lines.append("| " + " | ".join(table.headers) + " |")
            lines.append("| " + " | ".join(["---"] * table.columns) + " |")
            for row in table.data:
                lines.append("| " + " | ".join(row) + " |")
            lines.append("")

        return "\n".join(lines)

    def extract_text(self, file_path: str) -> str:
        """Extrahiert nur den Text aus einem Dokument"""
        result = self.analyze_document(file_path)
        return result.full_text if result else ""

    def extract_tables(self, file_path: str) -> List[TableData]:
        """Extrahiert nur die Tabellen aus einem Dokument"""
        result = self.analyze_document(file_path)
        return result.structure.tables if result else []

    def get_table_of_contents(self, file_path: str) -> List[Dict[str, Any]]:
        """Extrahiert das Inhaltsverzeichnis"""
        result = self.analyze_document(file_path)
        return result.structure.toc if result else []

    # =========================================================================
    # UTILITY
    # =========================================================================

    def get_capabilities(self) -> Dict[str, bool]:
        """Gibt die verfuegbaren Faehigkeiten zurueck"""
        return {
            "pdf_parsing": self.has_pypdf,
            "pdf_tables": self.has_pdfplumber,
            "docx_parsing": self.has_docx,
            "markdown_parsing": True,  # Immer verfuegbar
            "text_parsing": True,
            "markdown_conversion": True,
        }


# =============================================================================
# STANDALONE TEST
# =============================================================================

if __name__ == "__main__":
    logging.basicConfig(level=logging.INFO)

    print("=" * 60)
    print("HOLO DOCUMENT - Test")
    print("=" * 60)

    doc = HoloDocument()

    print(f"\nVerfuegbare Features:")
    for cap, available in doc.get_capabilities().items():
        status = "ja" if available else "nein"
        print(f"  - {cap}: {status}")

    # Test mit Dokumenten (falls vorhanden)
    test_files = ["test.pdf", "test.docx", "README.md", "test.txt"]

    for test_file in test_files:
        if os.path.exists(test_file):
            print(f"\n--- Analysiere: {test_file} ---")

            result = doc.analyze_document(test_file)

            if result:
                print(f"Typ: {result.metadata.doc_type.value}")
                print(f"Titel: {result.metadata.title}")
                print(f"Woerter: {result.metadata.word_count}")
                print(f"Seiten: {result.metadata.page_count}")
                print(f"Struktur-Elemente: {len(result.structure.elements)}")
                print(f"Ueberschriften: {len(result.structure.headings)}")
                print(f"Tabellen: {len(result.structure.tables)}")

                if result.summary_sentences:
                    print(f"\nErste Saetze:")
                    for s in result.summary_sentences[:2]:
                        print(f"  - {s[:80]}...")

            break
    else:
        print("\nKeine Test-Dokumente gefunden.")
        print("Erstellen Sie PDF/DOCX/MD-Dateien fuer vollstaendigen Test.")

    print("\n" + "=" * 60)
    print("Test abgeschlossen!")
    print("=" * 60)
